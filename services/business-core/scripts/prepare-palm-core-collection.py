#!/usr/bin/env python3
"""Prepare Palm Living's Core Collection for a reviewed, repeatable import.

The script is deliberately offline: it reads one client-supplied workbook and
writes a JSON manifest plus extracted image files. It never connects to Wodana
or mutates the source workbook. Financial rates are generated only when the
operator supplies the explicit commercial approval flags.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import sys
import unicodedata
import uuid
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError as error:  # pragma: no cover - operator guidance
    raise SystemExit("openpyxl is required to prepare an Excel collection.") from error


SCHEMA_VERSION = 1
CATALOG_ID = "1b74d8a5-638c-5d81-9d39-1866f44dded7"
NAMESPACE = uuid.UUID("a89adf61-a847-4aed-bf24-131a1951e9d4")
IGNORED_SHEETS = {"Template", "Brochure"}
ENQUIRE = re.compile(r"\benquir(?:e|y)(?:\s+now)?\b", re.IGNORECASE)
SPACE = re.compile(r"\s+")
SAFE_FILE = re.compile(r"[^a-zA-Z0-9._-]+")


def compact(value: Any) -> str:
    if value is None:
        return ""
    return SPACE.sub(" ", unicodedata.normalize("NFKC", str(value)).strip())


def normalized(value: Any) -> str:
    return compact(value).casefold()


def stable_id(key: str) -> str:
    return str(uuid.uuid5(NAMESPACE, key))


def code(prefix: str, key: str) -> str:
    stem = re.sub(r"[^A-Z0-9]+", "_", unicodedata.normalize("NFKD", key).encode("ascii", "ignore").decode().upper()).strip("_")
    suffix = hashlib.sha256(key.encode()).hexdigest()[:8].upper()
    return f"{prefix}_{stem[: max(1, 70 - len(prefix))]}_{suffix}"[:80]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def price_value(value: Any) -> tuple[str, float | None, str]:
    if isinstance(value, bool):
        return "unpriced", None, compact(value)
    if isinstance(value, (int, float)) and float(value) > 0:
        return "fixed", round(float(value), 2), compact(value)
    text = compact(value)
    if ENQUIRE.search(text):
        return "on_request", None, text
    return "unpriced", None, text


def aggregate_mode(modes: list[str]) -> str:
    present = [mode for mode in modes if mode != "unpriced"]
    if present and all(mode == "fixed" for mode in present) and len(present) == len(modes):
        return "fixed"
    if present:
        return "on_request"
    return "unpriced"


def price_label(header: str) -> str:
    text = compact(header)
    text = re.sub(r"(?i)^price\s*\(\s*aed\s*\)\s*-?", "", text).strip(" -()")
    text = re.sub(r"(?i)^price\s*-?", "", text).strip(" -()")
    return text.title() if text else "Standard"


@dataclass
class SourceRow:
    sheet: str
    row: int
    category: str
    name: str
    remarks: str
    dimensions: dict[str, str]
    prices: list[dict[str, Any]]
    raw: list[str]


def provenance(source_digest: str, row: SourceRow, role: str, approved: bool, extra: str = "") -> dict[str, Any]:
    source_key = f"{row.sheet}/row/{row.row}/{role}{('/' + extra) if extra else ''}"
    fingerprint = hashlib.sha256(json.dumps({
        "sheet": row.sheet,
        "row": row.row,
        "category": row.category,
        "name": row.name,
        "remarks": row.remarks,
        "dimensions": row.dimensions,
        "prices": row.prices,
        "raw": row.raw,
        "role": role,
        "extra": extra,
    }, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()).hexdigest()
    return {
        "namespace": f"xlsx:palm-core-collection:{source_digest[:16]}",
        "sourceKey": source_key,
        "sourceFingerprint": fingerprint,
        "reviewStatus": "approved" if approved else "needs_review",
        "sourceFileName": "Palm Living Core Collection.xlsx",
        "sourceSheet": row.sheet,
        "sourceRow": row.row,
    }


def base_record(*, record_id: str, catalog_id: str, kind: str, record_code: str, name: str,
                description: str = "", attributes: list[dict[str, str]] | None = None,
                links: list[dict[str, str]] | None = None, status: str = "needs_specification",
                commerce: str | None = None, record_provenance: dict[str, Any] | None = None,
                configuration: dict[str, str | int | float] | None = None,
                aliases: list[str] | None = None, dimension: str = "none", base_factor: float | None = None) -> dict[str, Any]:
    record: dict[str, Any] = {
        "id": record_id,
        "catalogId": catalog_id,
        "version": 0,
        "kind": kind,
        "code": record_code,
        "name": name[:180],
        "description": description[:4000],
        "market": "AE",
        "aliases": (aliases or [])[:30],
        "attributes": (attributes or [])[:30],
        "sizes": [],
        "sources": [],
        "links": links or [],
        "specificationStatus": status,
        "dimension": dimension,
    }
    if base_factor is not None:
        record["baseFactor"] = base_factor
    if configuration:
        record["configuration"] = configuration
    if commerce:
        record["commerce"] = {"pricingMode": commerce}
    if record_provenance:
        record["provenance"] = record_provenance
    return record


def attributes_for(row: SourceRow, include_dimensions: bool = True) -> list[dict[str, str]]:
    attributes = [
        {"name": "Collection category", "value": row.category},
        {"name": "Source sheet", "value": row.sheet},
        {"name": "Source row", "value": str(row.row)},
    ]
    if row.remarks:
        attributes.append({"name": "Source remarks", "value": row.remarks[:1000]})
    if include_dimensions:
        attributes.extend({"name": name, "value": value[:1000]} for name, value in row.dimensions.items() if value)
    return attributes


def dimensions_label(row: SourceRow) -> str:
    present = [f"{key.removesuffix(' (CM)').title()} {value}" for key, value in row.dimensions.items() if value]
    return " · ".join(present)[:160] or f"Source row {row.row}"


def extract_rows(workbook_path: Path) -> tuple[Any, list[SourceRow], dict[tuple[str, int], list[tuple[Any, int]]]]:
    workbook = load_workbook(workbook_path, read_only=False, data_only=True)
    rows: list[SourceRow] = []
    images: dict[tuple[str, int], list[tuple[Any, int]]] = defaultdict(list)
    for sheet in workbook.worksheets:
        if sheet.sheet_state != "visible" or sheet.title in IGNORED_SHEETS:
            continue
        headers = [compact(sheet.cell(3, column).value) for column in range(1, min(sheet.max_column, 60) + 1)]
        price_columns = [index + 1 for index, header in enumerate(headers) if "PRICE" in header.upper()]
        if not price_columns:
            continue
        dimension_columns = {
            index + 1: header.upper()
            for index, header in enumerate(headers)
            if any(axis in header.upper() for axis in ("HEIGHT", "LENGTH", "DEPTH"))
        }
        for row_number in range(4, sheet.max_row + 1):
            category = compact(sheet.cell(row_number, 2).value)
            name = compact(sheet.cell(row_number, 3).value)
            if not category or not name:
                continue
            prices = []
            for column in price_columns:
                raw = sheet.cell(row_number, column).value
                mode, amount, raw_text = price_value(raw)
                if raw is None or not raw_text:
                    continue
                prices.append({
                    "column": column,
                    "label": price_label(headers[column - 1]),
                    "mode": mode,
                    "amount": amount,
                    "raw": raw_text,
                })
            dimensions = {
                header: compact(sheet.cell(row_number, column).value)
                for column, header in dimension_columns.items()
                if compact(sheet.cell(row_number, column).value)
            }
            raw = [compact(sheet.cell(row_number, column).value) for column in range(1, max(price_columns + list(dimension_columns) + [3]) + 1)]
            rows.append(SourceRow(sheet.title, row_number, category, name, compact(sheet.cell(row_number, 1).value), dimensions, prices, raw))
        for image in sheet._images:  # openpyxl exposes anchored workbook media here.
            anchor = getattr(image, "anchor", None)
            if not hasattr(anchor, "_from"):
                continue
            row_number = anchor._from.row + 1
            column_number = anchor._from.col + 1
            if row_number >= 4:
                images[(sheet.title, row_number)].append((image, column_number))
    return workbook, rows, images


def prepare(args: argparse.Namespace) -> dict[str, Any]:
    source = args.workbook.resolve()
    output = args.output.resolve()
    if not source.is_file() or source.suffix.lower() != ".xlsx":
        raise SystemExit("Choose an existing .xlsx Palm Living Core Collection workbook.")
    if output == source.parent or output == source:
        raise SystemExit("Choose a dedicated output directory; the source directory is not modified.")
    if output.exists():
        if not args.replace:
            raise SystemExit(f"Output already exists: {output}. Use --replace to recreate it.")
        shutil.rmtree(output)
    output.mkdir(parents=True)
    assets_dir = output / "assets"
    assets_dir.mkdir()

    source_digest = sha256(source)
    workbook, rows, image_map = extract_rows(source)
    if len(rows) != args.expected_rows:
        raise SystemExit(f"Expected {args.expected_rows} item rows; found {len(rows)}. Review the workbook before importing.")

    catalog_id = CATALOG_ID
    unit_id = stable_id("palm-core:unit:each")
    approved = bool(args.approve_selling_prices)
    status = "workspace_verified" if approved else "needs_specification"
    records: list[dict[str, Any]] = [base_record(
        record_id=unit_id,
        catalog_id=catalog_id,
        kind="unit",
        record_code="EACH",
        name="Each",
        description="One finished collection item.",
        aliases=["nos", "piece"],
        status="workspace_verified",
        dimension="count",
        base_factor=1,
    )]
    group_ids: dict[str, str] = {}
    for order, sheet in enumerate(dict.fromkeys(row.sheet for row in rows), start=1):
        group_id = stable_id(f"palm-core:group:{normalized(sheet)}")
        group_ids[sheet] = group_id
        records.append(base_record(
            record_id=group_id,
            catalog_id=catalog_id,
            kind="record_group",
            record_code=code("GROUP", sheet),
            name=sheet,
            description=f"Palm Living Core Collection · {sheet}",
            status="workspace_verified",
            configuration={"defaultKind": "work", "icon": "items", "sortOrder": order * 10},
        ))

    grouped: dict[tuple[str, str], list[SourceRow]] = defaultdict(list)
    for row in rows:
        grouped[(row.sheet, normalized(row.name))].append(row)
    candidate_names = [items[0].name for items in grouped.values()]
    global_names = Counter(normalized(name) for name in candidate_names)
    assets: list[dict[str, Any]] = []
    rate_count = 0
    work_count = 0
    variant_count = 0

    def add_rate(owner: dict[str, Any], amount: float, source_row: SourceRow, label: str) -> None:
        nonlocal rate_count
        if not approved:
            return
        rate_id = stable_id(f"palm-core:rate:{owner['id']}:{amount}:{args.effective_from}:{label}")
        rate_name = f"{owner['name']} · selling · {args.effective_from}"[:180]
        records.append(base_record(
            record_id=rate_id,
            catalog_id=catalog_id,
            kind="rate_card",
            record_code=code("SELL", f"{owner['code']}:{args.effective_from}:{label}"),
            name=rate_name,
            description="Client-supplied Palm Living Core Collection selling price.",
            links=[
                {"relation": owner["kind"], "targetId": owner["id"]},
                {"relation": "unit", "targetId": unit_id},
                {"relation": "record_group", "targetId": group_ids[source_row.sheet]},
            ],
            status="workspace_verified",
            configuration={
                "currency": "AED",
                "amount": amount,
                "effectiveFrom": args.effective_from,
                "costIncludes": args.cost_includes,
                "taxIncluded": args.tax_included,
                "reference": f"Palm Living Core Collection · {source_row.sheet} · row {source_row.row}",
                "priceType": "selling",
            },
        ))
        rate_count += 1

    def extract_assets(row: SourceRow, owner: dict[str, Any]) -> None:
        anchored = sorted(image_map.get((row.sheet, row.row), []), key=lambda pair: pair[1])
        primary_used = False
        for index, (image, column) in enumerate(anchored):
            data = image._data()
            digest = hashlib.sha256(data).hexdigest()
            extension = (image.format or "png").lower()
            filename = f"{SAFE_FILE.sub('-', row.sheet).strip('-').lower()}-{row.row}-{column}-{index}-{digest[:10]}.{extension}"
            target = assets_dir / filename
            target.write_bytes(data)
            if column >= 9:
                role = "technical_drawing"
            elif not primary_used:
                role = "primary_image"
                primary_used = True
            else:
                role = "gallery_image"
            asset_id = stable_id(f"palm-core:asset:{row.sheet}:{row.row}:{column}:{index}:{digest}")
            assets.append({
                "recordId": owner["id"],
                "assetId": asset_id,
                "role": role,
                "altText": f"{owner['name']} · {role.replace('_', ' ')}"[:300],
                "sortOrder": index,
                "path": f"assets/{filename}",
                "sha256": digest,
                "byteSize": len(data),
                "sourceSheet": row.sheet,
                "sourceRow": row.row,
                "sourceColumn": column,
            })

    for (sheet, _name_key), family in grouped.items():
        first = family[0]
        display_name = first.name
        aliases: list[str] = []
        if global_names[normalized(first.name)] > 1:
            display_name = f"{first.name} — {sheet}"
            aliases = [first.name]
        multi_price = any(len(row.prices) > 1 for row in family)
        variants_needed = len(family) > 1 or multi_price
        all_modes = [price["mode"] for row in family for price in row.prices] or ["unpriced"]
        work_id = stable_id(f"palm-core:work:{normalized(sheet)}:{normalized(first.name)}")
        work = base_record(
            record_id=work_id,
            catalog_id=catalog_id,
            kind="work",
            record_code=code("PLM", f"{sheet}:{first.name}"),
            name=display_name,
            description=f"Palm Living Core Collection · {first.category}",
            aliases=aliases,
            attributes=attributes_for(first, include_dimensions=not variants_needed),
            links=[
                {"relation": "unit", "targetId": unit_id},
                {"relation": "record_group", "targetId": group_ids[sheet]},
            ],
            status=status,
            commerce=aggregate_mode(all_modes),
            record_provenance=provenance(source_digest, first, "work", approved, f"family-{len(family)}"),
        )
        records.append(work)
        work_count += 1
        if not variants_needed:
            extract_assets(first, work)
            for price in first.prices:
                if price["mode"] == "fixed" and price["amount"] is not None:
                    add_rate(work, price["amount"], first, price["label"])
            continue

        for row_index, row in enumerate(family, start=1):
            price_entries = row.prices or [{"label": dimensions_label(row), "mode": "unpriced", "amount": None, "raw": ""}]
            for price_index, price in enumerate(price_entries, start=1):
                label = price["label"] if len(price_entries) > 1 else dimensions_label(row)
                if not label or label == "Standard":
                    label = f"Option {row_index}"
                variant_key = f"{sheet}:{row.row}:{price.get('column', 0)}:{price_index}:{label}"
                variant_id = stable_id(f"palm-core:variant:{variant_key}")
                variant = base_record(
                    record_id=variant_id,
                    catalog_id=catalog_id,
                    kind="variant",
                    record_code=code("PLMV", variant_key),
                    name=f"{display_name} · {label}"[:180],
                    description=f"{first.category} option from Palm Living Core Collection.",
                    attributes=attributes_for(row),
                    links=[
                        {"relation": "work", "targetId": work_id},
                        {"relation": "unit", "targetId": unit_id},
                        {"relation": "record_group", "targetId": group_ids[sheet]},
                    ],
                    status=status,
                    commerce=price["mode"],
                    record_provenance=provenance(source_digest, row, "variant", approved, f"{price.get('column', 0)}-{price_index}"),
                    configuration={"specificationTier": label},
                )
                records.append(variant)
                variant_count += 1
                if price_index == 1:
                    extract_assets(row, variant)
                if price["mode"] == "fixed" and price["amount"] is not None:
                    add_rate(variant, price["amount"], row, label)

    identity = Counter((record["kind"], record["code"]) for record in records)
    names = Counter((record["kind"], normalized(record["name"])) for record in records)
    duplicate_codes = [key for key, count in identity.items() if count > 1]
    duplicate_names = [key for key, count in names.items() if count > 1]
    if duplicate_codes or duplicate_names:
        raise SystemExit(f"Prepared catalog contains duplicate identities: codes={duplicate_codes[:5]} names={duplicate_names[:5]}")
    if any(asset["byteSize"] > 20 * 1024 * 1024 for asset in assets):
        raise SystemExit("A source image exceeds the 20 MB catalog upload limit.")

    manifest = {
        "schemaVersion": SCHEMA_VERSION,
        "preparedAt": date.today().isoformat(),
        "source": {
            "fileName": source.name,
            "sha256": source_digest,
            "worksheets": list(dict.fromkeys(row.sheet for row in rows)),
            "itemRows": len(rows),
        },
        "approval": {
            "sellingPricesApproved": approved,
            "effectiveFrom": args.effective_from if approved else None,
            "taxIncluded": args.tax_included if approved else None,
            "costIncludes": args.cost_includes if approved else None,
            "note": "No verified rate cards are generated until explicit commercial approval is supplied." if not approved else "Approved selling prices were generated from numeric AED cells only.",
        },
        "catalog": {
            "id": catalog_id,
            "name": "Palm Living Core Collection",
            "description": "Client-supplied finished furniture, decor and bespoke collection items with source-row provenance.",
        },
        "records": records,
        "assets": assets,
        "audit": {
            "groups": len(group_ids),
            "works": work_count,
            "variants": variant_count,
            "sellingRates": rate_count,
            "assets": len(assets),
            "fixedPriceCells": sum(price["mode"] == "fixed" for row in rows for price in row.prices),
            "onRequestCells": sum(price["mode"] == "on_request" for row in rows for price in row.prices),
            "unpricedRows": sum(not row.prices or all(price["mode"] == "unpriced" for price in row.prices) for row in rows),
            "records": len(records),
        },
    }
    manifest_path = output / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    (output / "SHA256SUMS").write_text(
        "".join(f"{sha256(path)}  {path.relative_to(output)}\n" for path in sorted(output.rglob("*")) if path.is_file() and path.name != "SHA256SUMS")
    )
    workbook.close()
    return manifest


def parser() -> argparse.ArgumentParser:
    result = argparse.ArgumentParser(description=__doc__)
    result.add_argument("workbook", type=Path)
    result.add_argument("output", type=Path)
    result.add_argument("--expected-rows", type=int, default=542)
    result.add_argument("--replace", action="store_true")
    result.add_argument("--approve-selling-prices", action="store_true", help="Generate verified AED selling rate cards from numeric source cells.")
    result.add_argument("--effective-from", default=date.today().isoformat())
    result.add_argument("--tax-included", choices=("yes", "no"))
    result.add_argument("--cost-includes", choices=("labour", "machine", "labour_and_machine", "all_inclusive", "supply_only"))
    return result


def main() -> None:
    args = parser().parse_args()
    if args.approve_selling_prices and (not args.tax_included or not args.cost_includes):
        raise SystemExit("Approved prices require --tax-included and --cost-includes; these values are never inferred.")
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", args.effective_from):
        raise SystemExit("Use YYYY-MM-DD for --effective-from.")
    manifest = prepare(args)
    print(json.dumps({"ok": True, "catalogId": manifest["catalog"]["id"], **manifest["audit"]}, indent=2))


if __name__ == "__main__":
    main()
